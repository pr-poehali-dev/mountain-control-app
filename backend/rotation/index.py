import json
import os
import base64
from io import BytesIO
from datetime import datetime, date, timedelta
import psycopg2
import schedule as sch

COLORS = ['amber', 'cyan', 'green', 'purple', 'red']
MOVE_STATUSES = {'planned', 'confirmed', 'tickets', 'arrived', 'departed', 'no_show', 'cancelled'}


def get_db():
    return psycopg2.connect(os.environ['DATABASE_URL'])


def ser(o):
    if isinstance(o, (datetime, date)):
        return o.isoformat()
    return str(o)


def resp(status, body):
    return {
        'statusCode': status,
        'headers': {
            'Content-Type': 'application/json',
            'Access-Control-Allow-Origin': '*',
            'Access-Control-Allow-Methods': 'GET, POST, PUT, DELETE, OPTIONS',
            'Access-Control-Allow-Headers': 'Content-Type, Authorization, X-Authorization, X-Demo',
        },
        'body': json.dumps(body, ensure_ascii=False, default=ser),
    }


def q(s):
    return str(s or '').replace("'", "''")


def to_date(s, default=None):
    try:
        return date.fromisoformat(str(s)[:10])
    except (TypeError, ValueError):
        return default


def caller(event):
    """Администратор без демо-режима"""
    headers = event.get('headers') or {}
    params = event.get('queryStringParameters') or {}
    if headers.get('X-Demo', headers.get('x-demo', '')) == 'true' or params.get('demo') == '1':
        return None
    token = ''
    for k in ('X-Authorization', 'x-authorization', 'Authorization', 'authorization'):
        v = headers.get(k, '')
        if v:
            token = v[7:] if v.startswith('Bearer ') else v
            break
    token = token or params.get('auth', '')
    if not token:
        return None
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""SELECT u.id, u.full_name FROM users u JOIN sessions s ON s.user_id = u.id
                   WHERE s.token = '%s' AND s.expires_at > NOW() AND u.is_active = TRUE AND u.role = 'admin'
                   AND COALESCE(s.is_demo, FALSE) = FALSE""" % q(token))
    row = cur.fetchone()
    cur.close()
    conn.close()
    return {'id': row[0], 'name': row[1]} if row else None


# ---------- загрузка данных ----------

def load_departments(cur, only_id=None):
    where = "WHERE is_active = TRUE" + (" AND id = %d" % int(only_id) if only_id else "")
    cur.execute("""SELECT id, name, code, aliases, on_days, off_days, handover_days, anchor_date, scheme, head_name, color, sort_order
                   FROM rotation_departments %s ORDER BY sort_order, id""" % where)
    deps = []
    for r in cur.fetchall():
        aliases = r[3] if isinstance(r[3], list) else json.loads(r[3] or '[]')
        deps.append({'id': r[0], 'name': r[1], 'code': r[2], 'aliases': aliases, 'on_days': r[4], 'off_days': r[5],
                     'handover_days': r[6], 'anchor_date': r[7], 'scheme': r[8], 'head_name': r[9], 'color': r[10],
                     'sort_order': r[11], 'crews': []})
    if not deps:
        return deps
    by_id = {d['id']: d for d in deps}
    cur.execute("SELECT id, department_id, name, offset_days, sort_order FROM rotation_crews WHERE department_id IN (%s) ORDER BY sort_order, id"
                % ','.join(str(d['id']) for d in deps))
    for r in cur.fetchall():
        by_id[r[1]]['crews'].append({'id': r[0], 'department_id': r[1], 'name': r[2], 'offset_days': r[3], 'members': []})
    return deps


def load_members(cur, deps):
    """Работники раздела (собственный список, не связан с «Персоналом»)"""
    crews = {c['id']: c for d in deps for c in d['crews']}
    if not crews:
        return {}
    cur.execute("""SELECT id, crew_id, is_key, notes, full_name, COALESCE(NULLIF(position, ''), 'Без должности'), tab_number, phone,
                          pass_valid_to, housing
                   FROM rotation_workers
                   WHERE crew_id IN (%s) AND is_active = TRUE
                   ORDER BY position, full_name""" % ','.join(str(i) for i in crews))
    people = {}
    for r in cur.fetchall():
        p = {'id': r[0], 'crew_id': r[1], 'is_key': r[2], 'notes': r[3], 'full_name': r[4], 'position': r[5], 'tab_number': r[6],
             'phone': r[7], 'pass_valid_to': r[8], 'housing': r[9], 'exceptions': []}
        crews[r[1]]['members'].append(p)
        people[r[0]] = p
    if people:
        cur.execute("""SELECT id, worker_id, kind, date_from, date_to, notes FROM rotation_absences
                       WHERE worker_id IN (%s) ORDER BY date_from""" % ','.join(str(i) for i in people))
        for r in cur.fetchall():
            people[r[1]]['exceptions'].append({'id': r[0], 'kind': r[2], 'date_from': r[3], 'date_to': r[4], 'notes': r[5],
                                               'label': sch.EXC_LABELS.get(r[2], r[2])})
    return people


def load_norms(cur, dep_ids):
    if not dep_ids:
        return {}
    cur.execute("SELECT id, department_id, position, min_count FROM rotation_norms WHERE department_id IN (%s) ORDER BY position"
                % ','.join(str(i) for i in dep_ids))
    res = {}
    for r in cur.fetchall():
        res.setdefault(r[1], []).append({'id': r[0], 'position': r[2], 'min_count': r[3]})
    return res


def unassigned(cur, dep):
    cur.execute("""SELECT id, full_name, COALESCE(NULLIF(position, ''), 'Без должности'), tab_number
                   FROM rotation_workers
                   WHERE department_id = %d AND crew_id IS NULL AND is_active = TRUE
                   ORDER BY 3, 2""" % int(dep['id']))
    return [{'id': r[0], 'full_name': r[1], 'position': r[2], 'tab_number': r[3]} for r in cur.fetchall()]


# ---------- расчёты ----------

def present(dep, crew, person, d):
    return sch.is_on(dep, crew, d) and not sch.exception_on(person['exceptions'], d)


def coverage_gaps(dep, norms, d_from, days):
    """Дни, когда на объекте не хватает людей по нормативу (с учётом отпусков, больничных, увольнений)"""
    gaps = []
    norm_map = {n['position']: n['min_count'] for n in norms}
    for i in range(days):
        d = d_from + timedelta(days=i)
        counts = {}
        total = 0
        for c in dep['crews']:
            if not sch.is_on(dep, c, d):
                continue
            for p in c['members']:
                if sch.exception_on(p['exceptions'], d):
                    continue
                counts[p['position']] = counts.get(p['position'], 0) + 1
                total += 1
        for pos, need in norm_map.items():
            have = counts.get(pos, 0)
            if have < need:
                gaps.append({'date': d, 'position': pos, 'need': need, 'have': have})
    merged = {}
    for g in gaps:
        merged.setdefault((g['position'], g['need'], g['have']), []).append(g['date'])
    out = []
    for (pos, need, have), ds in merged.items():
        for a, b in sch.merge_days(ds):
            out.append({'position': pos, 'need': need, 'have': have, 'date_from': a, 'date_to': b})
    out.sort(key=lambda x: (x['date_from'], x['position']))
    return out


def crew_summary(dep, crew, today):
    on = sch.is_on(dep, crew, today)
    nxt, direction = sch.next_change(dep, crew, today)
    p = sch.phase(dep, crew, today)
    day_no = p + 1 if on else p - dep['on_days'] + 1
    present_cnt = sum(1 for m in crew['members'] if present(dep, crew, m, today))
    return {
        'id': crew['id'], 'name': crew['name'], 'offset_days': crew['offset_days'], 'on_site': on,
        'day_no': day_no, 'period_days': dep['on_days'] if on else dep['off_days'],
        'next_change': nxt, 'next_direction': direction, 'days_to_change': (nxt - today).days,
        'members_count': len(crew['members']), 'present_count': present_cnt if on else 0,
    }


def readiness(person, arrive, period_end):
    """Готовность работника к заезду на вахту (по данным самого раздела)"""
    issues = []
    vt = person['pass_valid_to']
    if vt:
        if vt < arrive:
            issues.append({'code': 'pass_expired', 'text': 'Пропуск истечёт до заезда (%s)' % vt.strftime('%d.%m.%Y')})
        elif vt < period_end:
            issues.append({'code': 'pass_mid', 'text': 'Пропуск истечёт во время вахты (%s)' % vt.strftime('%d.%m.%Y')})
    if not person['housing']:
        issues.append({'code': 'housing', 'text': 'Не указано проживание'})
    for e in sch.overlaps(person['exceptions'], arrive, period_end):
        issues.append({'code': 'exception', 'text': '%s %s–%s' % (e['label'], e['date_from'].strftime('%d.%m'),
                                                                  e['date_to'].strftime('%d.%m') if e['kind'] != 'dismissal' else '…')})
    return issues


def load_moves(cur, ids, d_from, d_to):
    if not ids:
        return {}
    cur.execute("""SELECT worker_id, change_date, direction, status, transport, notes FROM rotation_worker_moves
                   WHERE worker_id IN (%s) AND change_date BETWEEN '%s' AND '%s'""" % (','.join(str(i) for i in ids), d_from, d_to))
    return {(r[0], r[1], r[2]): {'status': r[3], 'transport': r[4], 'notes': r[5]} for r in cur.fetchall()}


def build_changes(cur, deps, d_from, d_to, dep_filter=None):
    """Список перевахтовок с составами заезда/выезда и готовностью"""
    pids = [m['id'] for d in deps for c in d['crews'] for m in c['members']]
    moves = load_moves(cur, pids, d_from - timedelta(days=1), d_to + timedelta(days=1))
    today = date.today()
    result = []
    for dep in deps:
        if dep_filter and dep['id'] != dep_filter:
            continue
        crews = {c['id']: c for c in dep['crews']}
        for d, sides in sch.changes(dep, dep['crews'], d_from, d_to):
            entry = {'date': d, 'department_id': dep['id'], 'department': dep['name'], 'color': dep['color'],
                     'days_left': (d - today).days, 'in': [], 'out': [], 'issues': 0, 'done': 0, 'total': 0}
            for side in ('in', 'out'):
                for cid in sides[side]:
                    crew = crews[cid]
                    period_end = d + timedelta(days=dep['on_days'] - 1)
                    people = []
                    for m in crew['members']:
                        mv = moves.get((m['id'], d, side), {})
                        exc = sch.exception_on(m['exceptions'], d if side == 'in' else d - timedelta(days=1))
                        issues = readiness(m, d, period_end) if side == 'in' else []
                        status = mv.get('status') or ('cancelled' if exc and exc['kind'] == 'dismissal' else 'planned')
                        overdue = d < today and status in ('planned', 'confirmed', 'tickets')
                        people.append({
                            'id': m['id'], 'full_name': m['full_name'], 'position': m['position'], 'is_key': m['is_key'],
                            'phone': m['phone'], 'status': status, 'transport': mv.get('transport', ''), 'notes': mv.get('notes', ''),
                            'issues': issues, 'exception': exc['label'] if exc else '', 'overdue': overdue,
                            'residence': m['housing'], 'tab_number': m['tab_number'],
                        })
                        entry['total'] += 1
                        if status in ('arrived', 'departed', 'cancelled'):
                            entry['done'] += 1
                        if issues or overdue:
                            entry['issues'] += 1
                    entry[side].append({'crew_id': cid, 'crew': crew['name'], 'people': people})
            result.append(entry)
    result.sort(key=lambda e: (e['date'], e['department']))
    return result


# ---------- действия ----------

def overview(cur):
    today = date.today()
    deps = load_departments(cur)
    people = load_members(cur, deps)
    norms = load_norms(cur, [d['id'] for d in deps])
    cards = []
    alerts = []
    for dep in deps:
        crews = [crew_summary(dep, c, today) for c in dep['crews']]
        gaps = coverage_gaps(dep, norms.get(dep['id'], []), today, 120)
        un = unassigned(cur, dep)
        on_site = sum(c['present_count'] for c in crews)
        nxt = min((c['next_change'] for c in crews), default=None)
        cards.append({
            'id': dep['id'], 'name': dep['name'], 'code': dep['code'], 'color': dep['color'], 'head_name': dep['head_name'],
            'on_days': dep['on_days'], 'off_days': dep['off_days'], 'scheme': dep['scheme'], 'crews': crews,
            'members_total': sum(c['members_count'] for c in crews), 'on_site': on_site, 'unassigned': len(un),
            'next_change': nxt, 'days_to_change': (nxt - today).days if nxt else None, 'gaps': gaps[:20], 'gaps_count': len(gaps),
            'norms_count': len(norms.get(dep['id'], [])),
        })
        for g in gaps[:5]:
            alerts.append({'level': 'red', 'department_id': dep['id'], 'kind': 'coverage',
                           'text': '%s: не хватает «%s» — %d из %d, %s–%s' % (dep['name'], g['position'], g['have'], g['need'],
                                                                         g['date_from'].strftime('%d.%m'), g['date_to'].strftime('%d.%m'))})
        if un:
            alerts.append({'level': 'amber', 'department_id': dep['id'], 'kind': 'unassigned',
                           'text': '%s: %d чел. не распределены по вахтам — нажмите «Автосборка»' % (dep['name'], len(un))})
        if (not dep['crews'] or all(len(c['members']) == 0 for c in dep['crews'])) and not un:
            alerts.append({'level': 'amber', 'department_id': dep['id'], 'kind': 'empty',
                           'text': '%s: нет работников — добавьте или загрузите список' % dep['name']})
    changes = build_changes(cur, deps, today - timedelta(days=3), today + timedelta(days=30))
    for ch in changes:
        if ch['issues'] and ch['days_left'] <= 14:
            alerts.append({'level': 'red' if ch['days_left'] < 0 else 'amber', 'department_id': ch['department_id'], 'kind': 'change',
                           'date': ch['date'],
                           'text': '%s, перевахтовка %s: проблем у %d чел.' % (ch['department'], ch['date'].strftime('%d.%m'), ch['issues'])})
    totals = {
        'departments': len(deps), 'people': len(people),
        'on_site': sum(c['on_site'] for c in cards), 'unassigned': sum(c['unassigned'] for c in cards),
        'gaps': sum(c['gaps_count'] for c in cards),
        'changes_30': len([c for c in changes if c['days_left'] >= 0]),
    }
    upcoming = [{k: v for k, v in c.items() if k not in ('in', 'out')} | {
        'in_count': sum(len(x['people']) for x in c['in']), 'out_count': sum(len(x['people']) for x in c['out'])}
        for c in changes if c['days_left'] >= -3][:12]
    return resp(200, {'today': today, 'totals': totals, 'departments': cards, 'alerts': alerts, 'upcoming': upcoming})


def dept_save(cur, body):
    name = (body.get('name') or '').strip()
    if not name:
        return resp(400, {'error': 'Укажите название подразделения'})
    on_days = max(1, min(366, int(body.get('on_days') or 60)))
    off_days = max(0, min(366, int(body.get('off_days') or 60)))
    handover = max(0, min(14, int(body.get('handover_days') or 1)))
    anchor = to_date(body.get('anchor_date'), date.today())
    scheme = 'staggered' if body.get('scheme') == 'staggered' else 'full'
    aliases = [a.strip() for a in (body.get('aliases') or []) if isinstance(a, str) and a.strip()]
    color = body.get('color') if body.get('color') in COLORS else 'amber'
    fields = (q(name), q(body.get('code', '')), q(json.dumps(aliases, ensure_ascii=False)), on_days, off_days, handover, anchor, scheme,
              q(body.get('head_name', '')), color)
    did = body.get('id')
    if did:
        cur.execute("SELECT scheme, on_days, off_days FROM rotation_departments WHERE id = %d" % int(did))
        old = cur.fetchone()
        if not old:
            return resp(404, {'error': 'Подразделение не найдено'})
        cur.execute("""UPDATE rotation_departments SET name = '%s', code = '%s', aliases = '%s'::jsonb, on_days = %d, off_days = %d,
                       handover_days = %d, anchor_date = '%s', scheme = '%s', head_name = '%s', color = '%s', updated_at = NOW()
                       WHERE id = %d""" % (fields + (int(did),)))
        if (old[0], old[1], old[2]) != (scheme, on_days, off_days):
            rebuild_crews(cur, int(did), scheme, on_days, off_days)
    else:
        cur.execute("""INSERT INTO rotation_departments (name, code, aliases, on_days, off_days, handover_days, anchor_date, scheme, head_name, color, sort_order)
                       VALUES ('%s', '%s', '%s'::jsonb, %d, %d, %d, '%s', '%s', '%s', '%s',
                               (SELECT COALESCE(MAX(sort_order), 0) + 1 FROM rotation_departments)) RETURNING id""" % fields)
        did = cur.fetchone()[0]
        rebuild_crews(cur, did, scheme, on_days, off_days)
    return resp(200, {'id': did, 'message': 'Сохранено'})


def rebuild_crews(cur, dep_id, scheme, on_days, off_days):
    """Пересоздаёт вахты под схему; людей переносит в вахты с тем же «полюсом» (А/Б)"""
    cur.execute("SELECT id, name FROM rotation_crews WHERE department_id = %d ORDER BY sort_order, id" % dep_id)
    old = cur.fetchall()
    templates = sch.crew_templates(scheme, on_days, off_days)
    keep = {}
    for i, (name, offset) in enumerate(templates):
        if i < len(old):
            cur.execute("UPDATE rotation_crews SET name = '%s', offset_days = %d, sort_order = %d WHERE id = %d" % (name, offset, i, old[i][0]))
            keep[name[0]] = keep.get(name[0], old[i][0])
        else:
            cur.execute("INSERT INTO rotation_crews (department_id, name, offset_days, sort_order) VALUES (%d, '%s', %d, %d) RETURNING id"
                        % (dep_id, name, offset, i))
            cid = cur.fetchone()[0]
            keep[name[0]] = keep.get(name[0], cid)
    for oid, oname in old[len(templates):]:
        target = keep.get(oname[0]) or list(keep.values())[0]
        cur.execute("UPDATE rotation_workers SET crew_id = %d WHERE crew_id = %d" % (target, oid))
        cur.execute("DELETE FROM rotation_crews WHERE id = %d" % oid)


def dept_delete(cur, body):
    did = int(body.get('id') or 0)
    cur.execute("UPDATE rotation_departments SET is_active = FALSE, updated_at = NOW() WHERE id = %d" % did)
    cur.execute("UPDATE rotation_workers SET crew_id = NULL WHERE department_id = %d" % did)
    return resp(200, {'message': 'Подразделение удалено из перевахтовки'})


def dept_detail(cur, params):
    did = int(params.get('id') or 0)
    deps = load_departments(cur, did)
    if not deps:
        return resp(404, {'error': 'Подразделение не найдено'})
    dep = deps[0]
    load_members(cur, deps)
    today = date.today()
    norms = load_norms(cur, [did]).get(did, [])
    positions = {}
    for c in dep['crews']:
        for m in c['members']:
            positions.setdefault(m['position'], {}).setdefault(c['id'], 0)
            positions[m['position']][c['id']] += 1
    crews = []
    for c in dep['crews']:
        s = crew_summary(dep, c, today)
        s['members'] = [{k: v for k, v in m.items() if k not in ('crew_id',)} for m in c['members']]
        s['periods'] = sch.periods(dep, c, today, today + timedelta(days=365))[:4]
        crews.append(s)
    dep_out = {k: v for k, v in dep.items() if k != 'crews'}
    return resp(200, {
        'department': dep_out, 'crews': crews, 'norms': norms, 'unassigned': unassigned(cur, dep),
        'positions': [{'position': p, 'by_crew': v} for p, v in sorted(positions.items())],
        'gaps': coverage_gaps(dep, norms, today, 180),
    })


def crews_total_ratio(dep):
    """Во сколько раз весь штат больше состава, одновременно находящегося на объекте"""
    cycle = sch.cycle_of(dep)
    return max(round(cycle / max(dep['on_days'], 1)), 1)


def auto_assign(cur, body):
    did = int(body.get('department_id') or 0)
    deps = load_departments(cur, did)
    if not deps or not deps[0]['crews']:
        return resp(404, {'error': 'Подразделение не найдено'})
    dep = deps[0]
    load_members(cur, deps)
    cand = unassigned(cur, dep)
    if body.get('rebalance'):
        for c in dep['crews']:
            cand += [{'id': m['id'], 'full_name': m['full_name'], 'position': m['position']} for m in c['members'] if not m['is_key']]
            c['members'] = [m for m in c['members'] if m['is_key']]
    crew_ids = [c['id'] for c in dep['crews']]
    plan = sch.balance({c['id']: c['members'] for c in dep['crews']}, cand, crew_ids)
    for wid, cid in plan:
        cur.execute("UPDATE rotation_workers SET crew_id = %d, updated_at = NOW() WHERE id = %d" % (cid, wid))
    cur.execute("""INSERT INTO rotation_norms (department_id, position, min_count)
                   SELECT %d, x.position, GREATEST(1, FLOOR(x.cnt::numeric / %d))::int FROM (
                       SELECT COALESCE(NULLIF(position, ''), 'Без должности') AS position, COUNT(*) AS cnt
                       FROM rotation_workers WHERE department_id = %d AND crew_id IS NOT NULL AND is_active = TRUE GROUP BY 1) x
                   ON CONFLICT (department_id, position) DO NOTHING""" % (did, crews_total_ratio(dep), did))
    return resp(200, {'assigned': len(plan), 'message': 'Распределено: %d чел.' % len(plan)})


def member_set(cur, body):
    """Перенос работника в вахту / снятие с вахты / закрепление"""
    wid = int(body.get('worker_id') or body.get('personnel_id') or 0)
    sets = ["updated_at = NOW()"]
    if 'crew_id' in body:
        crew = body.get('crew_id')
        sets.append("crew_id = %s" % ('NULL' if crew in (None, '', 0) else int(crew)))
    if body.get('is_key') is not None:
        sets.append("is_key = %s" % ('TRUE' if body.get('is_key') else 'FALSE'))
    cur.execute("UPDATE rotation_workers SET %s WHERE id = %d" % (', '.join(sets), wid))
    return resp(200, {'message': 'Сохранено'})


def worker_save(cur, body):
    """Добавление или изменение работника в списке подразделения"""
    name = ' '.join((body.get('full_name') or '').split())
    did = int(body.get('department_id') or 0)
    if not name or not did:
        return resp(400, {'error': 'Укажите ФИО и подразделение'})
    pv = to_date(body.get('pass_valid_to'))
    crew = body.get('crew_id')
    crew_sql = 'NULL' if crew in (None, '', 0, '0') else str(int(crew))
    vals = (q(name), q(body.get('position', '').strip()), q(body.get('tab_number', '').strip()), q(body.get('phone', '').strip()),
            ("'%s'" % pv) if pv else 'NULL', q(body.get('housing', '').strip()), q(body.get('notes', '')))
    if body.get('id'):
        cur.execute("""UPDATE rotation_workers SET full_name = '%s', position = '%s', tab_number = '%s', phone = '%s', pass_valid_to = %s,
                       housing = '%s', notes = '%s', crew_id = %s, updated_at = NOW() WHERE id = %d""" % (vals + (crew_sql, int(body['id']))))
        return resp(200, {'id': int(body['id']), 'message': 'Работник обновлён'})
    cur.execute("""INSERT INTO rotation_workers (full_name, position, tab_number, phone, pass_valid_to, housing, notes, crew_id, department_id)
                   VALUES ('%s', '%s', '%s', '%s', %s, '%s', '%s', %s, %d) RETURNING id""" % (vals + (crew_sql, did)))
    return resp(200, {'id': cur.fetchone()[0], 'message': 'Работник добавлен'})


def worker_delete(cur, body):
    cur.execute("UPDATE rotation_workers SET is_active = FALSE, crew_id = NULL, updated_at = NOW() WHERE id = %d" % int(body.get('id') or 0))
    return resp(200, {'message': 'Работник удалён из списка'})


def workers_import(cur, body):
    """Загрузка списка работников из Excel: ФИО, должность, табельный №, телефон, вахта, пропуск до, проживание"""
    from openpyxl import load_workbook
    did = int(body.get('department_id') or 0)
    deps = load_departments(cur, did)
    if not deps:
        return resp(404, {'error': 'Подразделение не найдено'})
    crews = {c['name'].strip().upper(): c['id'] for c in deps[0]['crews']}
    wb = load_workbook(BytesIO(base64.b64decode(body.get('file') or '')), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return resp(400, {'error': 'Файл пустой'})
    keys = {'фио': 'full_name', 'ф.и.о.': 'full_name', 'сотрудник': 'full_name', 'должность': 'position', 'профессия': 'position',
            'табельный': 'tab_number', 'таб': 'tab_number', 'телефон': 'phone', 'вахта': 'crew', 'пропуск': 'pass_valid_to',
            'проживание': 'housing', 'общежитие': 'housing', 'комната': 'housing'}
    header = [str(c or '').strip().lower() for c in rows[0]]
    cols = {}
    for idx, h in enumerate(header):
        for k, f in keys.items():
            if h.startswith(k) and f not in cols:
                cols[f] = idx
    data = rows[1:]
    if 'full_name' not in cols:
        cols = {'full_name': 0, 'position': 1, 'tab_number': 2, 'crew': 3}
        data = rows
    cur.execute("SELECT LOWER(full_name) FROM rotation_workers WHERE department_id = %d AND is_active = TRUE" % did)
    exist = {r[0] for r in cur.fetchall()}
    added = skipped = 0
    for r in data:
        get = lambda f: (r[cols[f]] if f in cols and cols[f] < len(r) else None)
        name = ' '.join(str(get('full_name') or '').split())
        if name in ('Иванов Иван Иванович', 'Петров Пётр Петрович') and str(get('tab_number') or '') in ('1001', '1002'):
            continue
        if not name or name.lower() in exist:
            skipped += 1 if name else 0
            continue
        pv = get('pass_valid_to')
        pv = pv.date() if isinstance(pv, datetime) else (pv if isinstance(pv, date) else to_date(pv) if pv else None)
        cname = str(get('crew') or '').strip().upper()
        crew = crews.get(cname)
        cur.execute("""INSERT INTO rotation_workers (department_id, crew_id, full_name, position, tab_number, phone, pass_valid_to, housing)
                       VALUES (%d, %s, '%s', '%s', '%s', '%s', %s, '%s')""" % (
            did, crew or 'NULL', q(name), q(str(get('position') or '').strip()), q(str(get('tab_number') or '').strip()),
            q(str(get('phone') or '').strip()), ("'%s'" % pv) if pv else 'NULL', q(str(get('housing') or '').strip())))
        exist.add(name.lower())
        added += 1
    return resp(200, {'added': added, 'skipped': skipped, 'message': 'Добавлено: %d, пропущено (уже есть): %d' % (added, skipped)})


def norm_save(cur, body):
    did = int(body.get('department_id') or 0)
    pos = (body.get('position') or '').strip()
    if not pos:
        return resp(400, {'error': 'Укажите должность'})
    cnt = max(0, int(body.get('min_count') or 0))
    if cnt == 0:
        cur.execute("DELETE FROM rotation_norms WHERE department_id = %d AND position = '%s'" % (did, q(pos)))
    else:
        cur.execute("""INSERT INTO rotation_norms (department_id, position, min_count) VALUES (%d, '%s', %d)
                       ON CONFLICT (department_id, position) DO UPDATE SET min_count = EXCLUDED.min_count""" % (did, q(pos), cnt))
    return resp(200, {'message': 'Норматив сохранён'})


def calendar(cur, params):
    d_from = to_date(params.get('from'), date.today().replace(day=1))
    months = max(1, min(12, int(params.get('months') or 6)))
    d_to = (d_from.replace(day=1) + timedelta(days=32 * months)).replace(day=1) - timedelta(days=1)
    dep_id = int(params.get('department_id') or 0) or None
    deps = load_departments(cur, dep_id)
    load_members(cur, deps)
    norms = load_norms(cur, [d['id'] for d in deps])
    rows = []
    for dep in deps:
        crews = []
        for c in dep['crews']:
            people = []
            if params.get('people') == '1':
                for m in c['members']:
                    absent = []
                    for e in sch.overlaps(m['exceptions'], d_from, d_to):
                        absent.append({'from': max(e['date_from'], d_from), 'to': d_to if e['kind'] == 'dismissal' else min(e['date_to'], d_to),
                                       'label': e['label']})
                    people.append({'id': m['id'], 'full_name': m['full_name'], 'position': m['position'], 'absent': absent})
            crews.append({'id': c['id'], 'name': c['name'], 'count': len(c['members']),
                          'periods': sch.periods(dep, c, d_from, d_to), 'people': people})
        gaps = coverage_gaps(dep, norms.get(dep['id'], []), d_from, (d_to - d_from).days + 1)
        rows.append({'id': dep['id'], 'name': dep['name'], 'color': dep['color'], 'crews': crews, 'gaps': gaps,
                     'changes': [c[0] for c in sch.changes(dep, dep['crews'], d_from, d_to)]})
    return resp(200, {'from': d_from, 'to': d_to, 'departments': rows})


def changes_list(cur, params):
    today = date.today()
    d_from = to_date(params.get('from'), today - timedelta(days=7))
    d_to = to_date(params.get('to'), today + timedelta(days=60))
    dep_id = int(params.get('department_id') or 0) or None
    deps = load_departments(cur)
    load_members(cur, deps)
    return resp(200, {'items': build_changes(cur, deps, d_from, d_to, dep_id)})


def move_set(cur, body, user):
    items = body.get('items') or [body]
    n = 0
    for it in items:
        wid = int(it.get('worker_id') or it.get('personnel_id') or 0)
        d = to_date(it.get('change_date'))
        direction = 'out' if it.get('direction') == 'out' else 'in'
        status = it.get('status') if it.get('status') in MOVE_STATUSES else 'planned'
        if not wid or not d:
            continue
        cur.execute("""INSERT INTO rotation_worker_moves (worker_id, change_date, direction, status, transport, notes)
                       VALUES (%d, '%s', '%s', '%s', '%s', '%s')
                       ON CONFLICT (worker_id, change_date, direction) DO UPDATE SET status = EXCLUDED.status,
                       transport = CASE WHEN EXCLUDED.transport <> '' THEN EXCLUDED.transport ELSE rotation_worker_moves.transport END,
                       notes = CASE WHEN EXCLUDED.notes <> '' THEN EXCLUDED.notes ELSE rotation_worker_moves.notes END, updated_at = NOW()""" % (
            wid, d, direction, status, q(it.get('transport', '')), q(it.get('notes', ''))))
        n += 1
    return resp(200, {'updated': n})


def exception_save(cur, body):
    wid = int(body.get('worker_id') or body.get('personnel_id') or 0)
    a = to_date(body.get('date_from'))
    b = to_date(body.get('date_to'), a)
    kind = body.get('kind') if body.get('kind') in sch.EXC_LABELS else 'other'
    if not wid or not a:
        return resp(400, {'error': 'Укажите работника и даты'})
    if b < a:
        a, b = b, a
    if body.get('id'):
        cur.execute("""UPDATE rotation_absences SET kind = '%s', date_from = '%s', date_to = '%s', notes = '%s' WHERE id = %d""" % (
            kind, a, b, q(body.get('notes', '')), int(body['id'])))
    else:
        cur.execute("""INSERT INTO rotation_absences (worker_id, kind, date_from, date_to, notes) VALUES (%d, '%s', '%s', '%s', '%s')""" % (
            wid, kind, a, b, q(body.get('notes', ''))))
    return resp(200, {'message': 'Сохранено'})


def exception_delete(cur, body):
    cur.execute("DELETE FROM rotation_absences WHERE id = %d" % int(body.get('id') or 0))
    return resp(200, {'message': 'Удалено'})


def exceptions_list(cur):
    cur.execute("""SELECT e.id, e.worker_id, w.full_name, w.position, COALESCE(d.name, ''), e.kind, e.date_from, e.date_to, e.notes
                   FROM rotation_absences e JOIN rotation_workers w ON w.id = e.worker_id
                   LEFT JOIN rotation_departments d ON d.id = w.department_id
                   WHERE w.is_active = TRUE AND (e.date_to >= CURRENT_DATE - 60 OR e.kind = 'dismissal')
                   ORDER BY e.date_from DESC LIMIT 500""")
    return resp(200, {'items': [{'id': r[0], 'personnel_id': r[1], 'full_name': r[2], 'position': r[3], 'department': r[4], 'kind': r[5],
                                 'label': sch.EXC_LABELS.get(r[5], r[5]), 'date_from': r[6], 'date_to': r[7], 'notes': r[8]}
                                for r in cur.fetchall()]})


def people_search(cur, params):
    s = (params.get('q') or '').strip()
    if len(s) < 2:
        return resp(200, {'items': []})
    cur.execute("""SELECT w.id, w.full_name, w.position, COALESCE(c.name, ''), d.name
                   FROM rotation_workers w JOIN rotation_departments d ON d.id = w.department_id AND d.is_active = TRUE
                   LEFT JOIN rotation_crews c ON c.id = w.crew_id
                   WHERE w.is_active = TRUE AND LOWER(w.full_name) LIKE '%%%s%%'
                   ORDER BY w.full_name LIMIT 20""" % q(s.lower()))
    return resp(200, {'items': [{'id': r[0], 'full_name': r[1], 'position': r[2], 'crew': r[3], 'rotation_department': r[4]}
                                for r in cur.fetchall()]})


def export_xlsx(cur, params):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    d_from = to_date(params.get('from'), date.today().replace(day=1))
    months = max(1, min(12, int(params.get('months') or 4)))
    d_to = (d_from.replace(day=1) + timedelta(days=32 * months)).replace(day=1) - timedelta(days=1)
    deps = load_departments(cur, int(params.get('department_id') or 0) or None)
    load_members(cur, deps)
    wb = Workbook()
    ws = wb.active
    ws.title = 'График вахт'
    days = (d_to - d_from).days + 1
    on_fill = PatternFill('solid', fgColor='F6B73C')
    exc_fill = PatternFill('solid', fgColor='EF6B6B')
    chg_fill = PatternFill('solid', fgColor='3FB6C8')
    head_fill = PatternFill('solid', fgColor='2A2F3A')
    thin = Side(style='thin', color='DDDDDD')
    ws.cell(1, 1, 'График вахт %s — %s (2/2: на вахте — жёлтый, отпуск/больничный — красный, день заезда — голубой)' % (
        d_from.strftime('%d.%m.%Y'), d_to.strftime('%d.%m.%Y'))).font = Font(bold=True)
    hdr = ['Подразделение', 'Вахта', 'ФИО', 'Должность']
    for i, h in enumerate(hdr, 1):
        c = ws.cell(3, i, h)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head_fill
    for i in range(days):
        d = d_from + timedelta(days=i)
        c = ws.cell(3, 5 + i, d.strftime('%d.%m'))
        c.font = Font(bold=True, color='FFFFFF', size=7)
        c.fill = head_fill
        c.alignment = Alignment(text_rotation=90, horizontal='center')
        ws.column_dimensions[get_column_letter(5 + i)].width = 3
    row = 4
    for dep in deps:
        for crew in dep['crews']:
            starts = {p[0] for p in sch.periods(dep, crew, d_from, d_to)}
            for m in crew['members']:
                ws.cell(row, 1, dep['name'])
                ws.cell(row, 2, crew['name'])
                ws.cell(row, 3, m['full_name'])
                ws.cell(row, 4, m['position'])
                for i in range(days):
                    d = d_from + timedelta(days=i)
                    cell = ws.cell(row, 5 + i)
                    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                    exc = sch.exception_on(m['exceptions'], d)
                    if exc:
                        cell.fill = exc_fill
                        cell.value = exc['label'][:1]
                    elif sch.is_on(dep, crew, d):
                        cell.fill = chg_fill if d in starts else on_fill
                row += 1
    for col, w in (('A', 22), ('B', 7), ('C', 32), ('D', 28)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'E4'
    ws2 = wb.create_sheet('Перевахтовки')
    ws2.append(['Дата', 'Подразделение', 'Направление', 'Вахта', 'ФИО', 'Должность', 'Статус', 'Транспорт', 'Замечания'])
    labels = {'planned': 'Запланирован', 'confirmed': 'Подтверждён', 'tickets': 'Билеты куплены', 'arrived': 'Прибыл',
              'departed': 'Выехал', 'no_show': 'Не явился', 'cancelled': 'Отменён'}
    for ch in build_changes(cur, deps, d_from, d_to):
        for side, title in (('in', 'Заезд'), ('out', 'Выезд')):
            for g in ch[side]:
                for p in g['people']:
                    ws2.append([ch['date'].strftime('%d.%m.%Y'), ch['department'], title, g['crew'], p['full_name'], p['position'],
                                labels.get(p['status'], p['status']), p['transport'], '; '.join(i['text'] for i in p['issues'])])
    for col, w in zip('ABCDEFGHI', (12, 22, 10, 7, 32, 26, 16, 18, 50)):
        ws2.column_dimensions[col].width = w
    buf = BytesIO()
    wb.save(buf)
    return resp(200, {'file': base64.b64encode(buf.getvalue()).decode(), 'file_name': 'Перевахтовка_%s.xlsx' % d_from.strftime('%Y_%m')})


def handler(event, context):
    """Перевахтовка: подразделения и вахты, автосборка составов, графики, план заездов/выездов, контроль укомплектованности"""
    if event.get('httpMethod') == 'OPTIONS':
        return resp(200, '')
    user = caller(event)
    if not user:
        return resp(403, {'error': 'Раздел «Перевахтовка» доступен только администратору'})
    method = event.get('httpMethod', 'GET')
    params = event.get('queryStringParameters') or {}
    action = params.get('action', '')
    body = json.loads(event.get('body') or '{}') if method != 'GET' else {}
    conn = get_db()
    cur = conn.cursor()
    routes = {
        ('GET', 'overview'): lambda: overview(cur),
        ('GET', 'department'): lambda: dept_detail(cur, params),
        ('POST', 'department'): lambda: dept_save(cur, body),
        ('POST', 'department-delete'): lambda: dept_delete(cur, body),
        ('POST', 'auto-assign'): lambda: auto_assign(cur, body),
        ('PUT', 'member'): lambda: member_set(cur, body),
        ('POST', 'worker'): lambda: worker_save(cur, body),
        ('POST', 'worker-delete'): lambda: worker_delete(cur, body),
        ('POST', 'workers-import'): lambda: workers_import(cur, body),
        ('PUT', 'norm'): lambda: norm_save(cur, body),
        ('GET', 'calendar'): lambda: calendar(cur, params),
        ('GET', 'changes'): lambda: changes_list(cur, params),
        ('PUT', 'move'): lambda: move_set(cur, body, user),
        ('GET', 'exceptions'): lambda: exceptions_list(cur),
        ('POST', 'exception'): lambda: exception_save(cur, body),
        ('POST', 'exception-delete'): lambda: exception_delete(cur, body),
        ('GET', 'people'): lambda: people_search(cur, params),
        ('GET', 'export'): lambda: export_xlsx(cur, params),
    }
    fn = routes.get((method, action))
    if not fn:
        cur.close()
        conn.close()
        return resp(400, {'error': 'Неизвестное действие: %s' % action})
    res = fn()
    if method != 'GET':
        conn.commit()
    cur.close()
    conn.close()
    return res
