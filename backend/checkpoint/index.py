import json
import os
from datetime import datetime, date as date_type
import psycopg2


def caller_is_admin(event):
    """Цифровой код сотрудников видит только администратор (не в демо-режиме)"""
    headers = event.get('headers') or {}
    params = event.get('queryStringParameters') or {}
    if headers.get('X-Demo', headers.get('x-demo', '')) == 'true' or params.get('demo') == '1':
        return False
    token = ''
    for key in ['X-Authorization', 'x-authorization', 'Authorization', 'authorization']:
        val = headers.get(key, '')
        if val:
            token = val[7:] if val.startswith('Bearer ') else val
            break
    if not token:
        token = params.get('auth', '')
    if not token:
        return False
    conn = psycopg2.connect(os.environ['DATABASE_URL'])
    cur = conn.cursor()
    cur.execute("""
        SELECT 1 FROM users u JOIN sessions s ON s.user_id = u.id
        WHERE s.token = '%s' AND s.expires_at > NOW() AND u.is_active = TRUE AND u.role = 'admin'
    """ % token.replace("'", "''"))
    ok = cur.fetchone() is not None
    cur.close()
    conn.close()
    return ok

def is_demo_request(event):
    headers = event.get('headers') or {}
    return headers.get('X-Demo', headers.get('x-demo', '')) == 'true'

def get_db():
    return psycopg2.connect(os.environ['DATABASE_URL'])

def serialize_default(obj):
    if isinstance(obj, datetime):
        if obj.tzinfo is None:
            return obj.isoformat() + '+00:00'
        return obj.isoformat()
    if isinstance(obj, (date_type,)):
        return obj.isoformat()
    return str(obj)


CODE_KEYS = ('personal_code', 'person_code')
HIDE_CODES = {'on': False}


def strip_codes(obj):
    if isinstance(obj, dict):
        return {k: ('' if k in CODE_KEYS else strip_codes(v)) for k, v in obj.items()}
    if isinstance(obj, list):
        return [strip_codes(v) for v in obj]
    return obj

def json_response(status, body):
    if HIDE_CODES['on']:
        body = strip_codes(body)
    return {
        'statusCode': status,
        'headers': {
            'Content-Type': 'application/json',
            'Access-Control-Allow-Origin': '*',
            'Access-Control-Allow-Methods': 'GET, POST, PUT, DELETE, OPTIONS',
            'Access-Control-Allow-Headers': 'Content-Type, Authorization, X-Authorization, X-Demo'
        },
        'body': json.dumps(body, ensure_ascii=False, default=serialize_default)
    }


def sync_rooms(cur):
    """Пересчитывает занятость и статус комнат по проживающим на предприятии"""
    cur.execute("""
        UPDATE rooms r SET occupied = s.cnt,
            status = CASE WHEN s.cnt = 0 THEN 'free' WHEN s.cnt >= r.capacity THEN 'full' ELSE 'partial' END
        FROM (SELECT r2.id, COUNT(p.id) AS cnt FROM rooms r2
              LEFT JOIN personnel p ON p.residence_room_id = r2.id
                AND p.status IN ('arrived', 'on_shift') AND COALESCE(p.is_hidden, FALSE) = FALSE
              GROUP BY r2.id) s
        WHERE s.id = r.id
    """)

DENY_SQL = "COALESCE(%s.notes, '') NOT LIKE 'ОТКАЗ%%'"

DEFAULT_POINT_SETTINGS = {
    'counts_territory': False,
    'require_on_site': True,
    'check_pass': True,
    'require_medical': False,
    'vehicles': False,
}

POINT_ICONS = ['DoorOpen', 'Factory', 'Warehouse', 'Pickaxe', 'Building2', 'Truck', 'Fuel', 'Zap', 'FlaskConical', 'ShieldCheck', 'HardHat', 'Mountain']


def territory_sql(alias='cp'):
    """Проходы, которые меняют статус «на территории»: главный КПП и объекты-ворота"""
    return ("(%s.control_point_id IS NULL OR %s.control_point_id IN (SELECT id FROM control_points "
            "WHERE is_main = TRUE OR (settings->>'counts_territory') = 'true'))") % (alias, alias)


def row_to_point(r):
    settings = r[8] if isinstance(r[8], dict) else json.loads(r[8] or '{}')
    merged = dict(DEFAULT_POINT_SETTINGS)
    if r[5]:
        merged.update({'counts_territory': True, 'require_on_site': False, 'vehicles': True})
    merged.update(settings or {})
    return {
        'id': r[0], 'name': r[1], 'short_name': r[2] or r[1], 'description': r[3] or '', 'icon': r[4] or 'DoorOpen',
        'is_main': r[5], 'is_active': r[6], 'sort_order': r[7], 'settings': merged,
        'page_key': 'checkpoint' if r[5] else 'cp_%d' % r[0],
    }


POINT_COLS = "id, name, short_name, description, icon, is_main, is_active, sort_order, settings"


def load_point(cur, point_id):
    """Объект контроля по id, без id — главный КПП"""
    pid = str(point_id or '').strip()
    if pid and pid.isdigit():
        cur.execute("SELECT %s FROM control_points WHERE id = %d AND is_active = TRUE" % (POINT_COLS, int(pid)))
    else:
        cur.execute("SELECT %s FROM control_points WHERE is_main = TRUE ORDER BY id LIMIT 1" % POINT_COLS)
    r = cur.fetchone()
    return row_to_point(r) if r else None


def point_sql(point, alias='cp'):
    """Фильтр проходов по конкретному объекту"""
    if not point or point['is_main']:
        return "(%s.control_point_id IS NULL OR %s.control_point_id = %d)" % (alias, alias, point['id'] if point else 0)
    return "%s.control_point_id = %d" % (alias, point['id'])


def is_territory_point(point):
    return not point or point['is_main'] or bool(point['settings'].get('counts_territory'))


def inside_sql(point, demo_val):
    """SQL со списком personnel_id, кто сейчас находится внутри объекта"""
    if is_territory_point(point):
        flt = territory_sql('cp')
        extra = ''
    else:
        flt = point_sql(point, 'cp')
        extra = (" AND NOT EXISTS (SELECT 1 FROM checkpoint_passes t WHERE t.personnel_id = ls.personnel_id "
                 "AND t.direction = 'out' AND %s AND %s AND t.created_at > ls.created_at)") % (territory_sql('t'), DENY_SQL % 't')
    return """
        SELECT ls.personnel_id, ls.created_at, ls.checkpoint_name FROM (
            SELECT DISTINCT ON (cp.personnel_id) cp.personnel_id, cp.direction, cp.created_at, cp.checkpoint_name
            FROM checkpoint_passes cp
            JOIN personnel pp ON pp.id = cp.personnel_id AND pp.is_demo_data = %s AND pp.status <> 'archived'
            WHERE cp.personnel_id IS NOT NULL AND %s AND %s
            ORDER BY cp.personnel_id, cp.created_at DESC, cp.id DESC
        ) ls WHERE ls.direction = 'in'%s
    """ % (demo_val, DENY_SQL % 'cp', flt, extra)


def caller_user(cur, event):
    token = get_auth_token(event)
    if not token:
        return None
    cur.execute("""
        SELECT u.id, u.full_name, u.role FROM users u JOIN sessions s ON s.user_id = u.id
        WHERE s.token = '%s' AND s.expires_at > NOW() AND u.is_active = TRUE
    """ % token.replace("'", "''"))
    return cur.fetchone()


def list_points(params):
    conn = get_db()
    cur = conn.cursor()
    show_all = params.get('all') == '1'
    cur.execute("SELECT %s FROM control_points %s ORDER BY is_main DESC, sort_order, id"
                % (POINT_COLS, '' if show_all else 'WHERE is_active = TRUE'))
    items = [row_to_point(r) for r in cur.fetchall()]
    cur.close()
    conn.close()
    return json_response(200, {'points': items, 'icons': POINT_ICONS})


def clean_settings(raw, is_main):
    raw = raw if isinstance(raw, dict) else {}
    out = {}
    for k in DEFAULT_POINT_SETTINGS:
        if k in raw:
            out[k] = bool(raw[k])
    if is_main:
        out['counts_territory'] = True
    return out


def save_point(event, body):
    conn = get_db()
    cur = conn.cursor()
    admin = get_admin(cur, event)
    if not admin:
        cur.close()
        conn.close()
        return json_response(403, {'error': 'Объекты контроля настраивает только администратор'})
    name = str(body.get('name') or '').strip()[:100]
    if not name:
        cur.close()
        conn.close()
        return json_response(400, {'error': 'Укажите название объекта'})
    short = str(body.get('short_name') or '').strip()[:30]
    desc = str(body.get('description') or '').strip()[:500]
    icon = str(body.get('icon') or 'DoorOpen')
    if icon not in POINT_ICONS:
        icon = 'DoorOpen'
    pid = body.get('id')
    q = lambda v: v.replace("'", "''")
    if pid:
        cur.execute("SELECT is_main, settings FROM control_points WHERE id = %d" % int(pid))
        ex = cur.fetchone()
        if not ex:
            cur.close()
            conn.close()
            return json_response(404, {'error': 'Объект не найден'})
        cur_set = ex[1] if isinstance(ex[1], dict) else json.loads(ex[1] or '{}')
        cur_set.update(clean_settings(body.get('settings'), ex[0]))
        st = json.dumps(cur_set, ensure_ascii=False)
        cur.execute("""UPDATE control_points SET name = '%s', short_name = '%s', description = '%s', icon = '%s',
                       settings = '%s'::jsonb, updated_at = NOW() WHERE id = %d"""
                    % (q(name), q(short), q(desc), icon, q(st), int(pid)))
        new_id = int(pid)
        cur.execute("INSERT INTO events (event_type, description) VALUES ('control_point', '%s')"
                    % q('Изменён объект контроля «%s»' % name))
    else:
        st = dict(DEFAULT_POINT_SETTINGS)
        st.update(clean_settings(body.get('settings'), False))
        cur.execute("SELECT COALESCE(MAX(sort_order), 0) + 1 FROM control_points")
        order = cur.fetchone()[0]
        cur.execute("""INSERT INTO control_points (name, short_name, description, icon, sort_order, settings)
                       VALUES ('%s', '%s', '%s', '%s', %d, '%s'::jsonb) RETURNING id"""
                    % (q(name), q(short), q(desc), icon, order, q(json.dumps(st, ensure_ascii=False))))
        new_id = cur.fetchone()[0]
        cur.execute("INSERT INTO events (event_type, description) VALUES ('control_point', '%s')"
                    % q('Добавлен объект контроля «%s»' % name))
    conn.commit()
    point = load_point(cur, new_id)
    cur.close()
    conn.close()
    return json_response(200, {'point': point})


def delete_point(event, body):
    conn = get_db()
    cur = conn.cursor()
    admin = get_admin(cur, event)
    if not admin:
        cur.close()
        conn.close()
        return json_response(403, {'error': 'Удалять объекты может только администратор'})
    pid = int(body.get('id') or 0)
    cur.execute("SELECT is_main, name FROM control_points WHERE id = %d" % pid)
    r = cur.fetchone()
    if not r:
        cur.close()
        conn.close()
        return json_response(404, {'error': 'Объект не найден'})
    if r[0]:
        cur.close()
        conn.close()
        return json_response(400, {'error': 'Главный КПП удалить нельзя'})
    cur.execute("UPDATE control_points SET is_active = FALSE, updated_at = NOW() WHERE id = %d" % pid)
    cur.execute("INSERT INTO events (event_type, description) VALUES ('control_point', '%s')"
                % ('Удалён объект контроля «%s»' % r[1]).replace("'", "''"))
    conn.commit()
    cur.close()
    conn.close()
    return json_response(200, {'success': True})


def points_summary(event):
    """Сводка по всем объектам: кто внутри, входы/выходы/отказы за сегодня"""
    demo_val = 'TRUE' if is_demo_request(event) else 'FALSE'
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT %s FROM control_points WHERE is_active = TRUE ORDER BY is_main DESC, sort_order, id" % POINT_COLS)
    points = [row_to_point(r) for r in cur.fetchall()]
    out = []
    for pt in points:
        cur.execute("SELECT COUNT(*) FROM (%s) x" % inside_sql(pt, demo_val))
        inside = cur.fetchone()[0]
        cur.execute("""
            SELECT COUNT(*) FILTER (WHERE cp.direction = 'in' AND %s),
                   COUNT(*) FILTER (WHERE cp.direction = 'out' AND %s),
                   COUNT(*) FILTER (WHERE NOT (%s)),
                   MAX(cp.created_at)
            FROM checkpoint_passes cp JOIN personnel p ON p.id = cp.personnel_id AND p.is_demo_data = %s
            WHERE %s AND cp.created_at::date = CURRENT_DATE
        """ % (DENY_SQL % 'cp', DENY_SQL % 'cp', DENY_SQL % 'cp', demo_val, point_sql(pt, 'cp')))
        t = cur.fetchone()
        item = dict(pt)
        item.update({'inside': inside, 'today_in': t[0], 'today_out': t[1], 'today_denied': t[2], 'last_pass_at': t[3]})
        out.append(item)
    cur.close()
    conn.close()
    return json_response(200, {'points': out})


def handler(event, context):
    """КПП и объекты контроля — вход/выход через сканер, журнал проходов, настройки объектов, связь с АХО"""
    if event.get('httpMethod') == 'OPTIONS':
        return json_response(200, '')
    HIDE_CODES['on'] = not caller_is_admin(event)

    method = event.get('httpMethod', 'GET')
    params = event.get('queryStringParameters') or {}
    action = params.get('action', '')
    body = json.loads(event.get('body', '{}') or '{}')

    if method == 'GET' and action == 'points':
        return list_points(params)
    elif method == 'POST' and action == 'point-save':
        return save_point(event, body)
    elif method == 'DELETE' and action == 'point-delete':
        return delete_point(event, body)
    elif method == 'GET' and action == 'points-summary':
        return points_summary(event)
    elif method == 'POST' and action == 'pass':
        return register_pass(body, event)
    elif method == 'GET' and action == 'journal':
        return get_journal(params, event)
    elif method == 'GET' and action == 'stats':
        return get_stats(event, params)
    elif method == 'GET' and action == 'on-site':
        return get_on_site(params, event)
    elif method == 'GET' and action == 'export':
        return export_journal(params, event)
    elif method == 'PUT' and action == 'pass-update':
        return update_pass(event, body)
    elif method == 'DELETE' and action == 'pass-delete':
        return delete_pass(event, body, params)
    elif method == 'GET' and action == 'vehicle-search':
        return vehicle_search(params, event)
    elif method == 'POST' and action == 'vehicle-pass':
        return vehicle_pass(body, event)
    elif method == 'GET' and action == 'vehicle-journal':
        return vehicle_journal(params, event)
    elif method == 'GET' and action == 'vehicle-export':
        return vehicle_export(params, event)

    return json_response(404, {'error': 'Маршрут не найден'})

def fmt_date(d):
    return d.strftime('%d.%m.%Y') if d else ''


def pass_period(permanent, d_from, d_to):
    if permanent:
        return 'Бессрочно'
    if d_from and d_to:
        return 'с %s по %s' % (fmt_date(d_from), fmt_date(d_to))
    if d_from:
        return 'с %s' % fmt_date(d_from)
    if d_to:
        return 'по %s' % fmt_date(d_to)
    return ''


def check_pass(number, permanent, d_from, d_to, title):
    """Возвращает состояние пропуска и причину отказа"""
    today = date_type.today()
    if permanent:
        return 'active', ''
    if not (number or '').strip() and not d_from and not d_to:
        return 'none', '%s не оформлен' % title
    if d_to and d_to < today:
        return 'expired', '%s истёк %s' % (title, fmt_date(d_to))
    if d_from and d_from > today:
        return 'future', '%s начнёт действовать с %s' % (title, fmt_date(d_from))
    if not d_from and not d_to:
        return 'none', 'У пропуска не указан срок действия'
    return 'active', ''


LAT = 'ABEKMHOPCTYX'
CYR = 'АВЕКМНОРСТУХ'


def norm_plate(val):
    v = ''.join(str(val or '').upper().split()).replace('-', '')
    return v.translate(str.maketrans(LAT, CYR))


PLATE_SQL = "TRANSLATE(REPLACE(REPLACE(UPPER(p.vehicle_number), ' ', ''), '-', ''), '%s', '%s')" % (LAT, CYR)


def find_vehicle(cur, plate, event):
    demo_val = 'TRUE' if is_demo_request(event) else 'FALSE'
    cur.execute("""
        SELECT p.id, p.full_name, p.position, p.organization, COALESCE(p.photo_url, ''),
               p.vehicle_number, p.vehicle_brand, p.vehicle_pass_number,
               p.vehicle_pass_permanent, p.vehicle_pass_valid_from, p.vehicle_pass_valid_to, p.personal_code
        FROM personnel p
        WHERE %s = '%s' AND p.status != 'archived' AND p.is_demo_data = %s
        ORDER BY p.id DESC LIMIT 1
    """ % (PLATE_SQL, plate.replace("'", "''"), demo_val))
    r = cur.fetchone()
    if not r:
        return None
    state, reason = check_pass(r[7], r[8], r[9], r[10], 'Пропуск на автомобиль')
    return {
        'personnel_id': r[0], 'owner_name': r[1], 'position': r[2] or '', 'organization': r[3] or '',
        'photo_url': r[4], 'vehicle_number': r[5], 'vehicle_brand': r[6] or '',
        'pass_number': r[7] or '', 'pass_period': pass_period(r[8], r[9], r[10]),
        'state': state, 'reason': reason, 'personal_code': r[11]
    }


def last_vehicle_direction(cur, plate):
    cur.execute("""
        SELECT direction FROM vehicle_passes
        WHERE TRANSLATE(REPLACE(REPLACE(UPPER(vehicle_number), ' ', ''), '-', ''), '%s', '%s') = '%s' AND allowed = TRUE
        ORDER BY created_at DESC, id DESC LIMIT 1
    """ % (LAT, CYR, plate.replace("'", "''")))
    r = cur.fetchone()
    return r[0] if r else ''


def vehicle_search(params, event):
    plate = norm_plate(params.get('number', ''))
    if len(plate) < 3:
        return json_response(400, {'error': 'Введите госномер полностью'})
    conn = get_db()
    cur = conn.cursor()
    v = find_vehicle(cur, plate, event)
    if v:
        v['on_site'] = last_vehicle_direction(cur, plate) == 'in'
    cur.close()
    conn.close()
    if not v:
        return json_response(404, {'error': 'Автомобиль %s не найден в базе пропусков' % plate})
    return json_response(200, {'vehicle': v})


def vehicle_pass(body, event):
    plate = norm_plate(body.get('number', ''))
    direction = body.get('direction', 'in')
    checkpoint_name = str(body.get('checkpoint_name') or 'КПП-1')[:100]
    if len(plate) < 3:
        return json_response(400, {'error': 'Введите госномер полностью'})
    if direction not in ('in', 'out'):
        return json_response(400, {'error': 'direction должен быть in или out'})
    conn = get_db()
    cur = conn.cursor()
    point = load_point(cur, body.get('control_point_id'))
    if point:
        checkpoint_name = point['short_name'][:100]
    v = find_vehicle(cur, plate, event)
    if v:
        allowed = direction == 'out' or v['state'] == 'active'
        reason = '' if allowed else v['reason']
    else:
        allowed = False
        reason = 'Автомобиль не найден в базе пропусков'
    if direction == 'out' and not v:
        allowed = True
        reason = ''
    notes = '' if allowed else 'ОТКАЗ: ' + reason
    cur.execute("""
        INSERT INTO vehicle_passes (personnel_id, vehicle_number, vehicle_brand, owner_name, direction, checkpoint_name, allowed, notes, control_point_id)
        VALUES (%s, '%s', '%s', '%s', '%s', '%s', %s, '%s', %s) RETURNING id, created_at
    """ % (v['personnel_id'] if v else 'NULL', plate, (v['vehicle_brand'] if v else '').replace("'", "''"),
           (v['owner_name'] if v else '').replace("'", "''"), direction, checkpoint_name.replace("'", "''"),
           'TRUE' if allowed else 'FALSE', notes.replace("'", "''"), point['id'] if point else 'NULL'))
    rec = cur.fetchone()
    if v:
        desc = '%s: автомобиль %s (%s) — %s' % (checkpoint_name, plate, v['owner_name'],
               ('въезд' if direction == 'in' else 'выезд') if allowed else 'отказ во въезде. ' + reason)
        cur.execute("INSERT INTO events (event_type, description, personnel_id) VALUES ('%s', '%s', %d)" % (
            'vehicle_' + (direction if allowed else 'denied'), desc.replace("'", "''"), v['personnel_id']))
    conn.commit()
    cur.close()
    conn.close()
    return json_response(200, {
        'result': 'allowed' if allowed else 'denied',
        'message': ('Въезд разрешён' if direction == 'in' else 'Выезд зафиксирован') if allowed else 'Въезд запрещён',
        'reason': reason,
        'direction': direction,
        'vehicle': v,
        'vehicle_number': plate,
        'pass_id': rec[0]
    })


def vehicle_journal(params, event):
    date_from = params.get('date_from', '')
    date_to = params.get('date_to', '')
    page = max(1, int(params.get('page', '1')))
    per_page = 50
    where = "WHERE 1=1"
    if date_from:
        where += " AND vp.created_at >= '%s 00:00:00'" % date_from.replace("'", "''")
    if date_to:
        where += " AND vp.created_at <= '%s 23:59:59'" % date_to.replace("'", "''")
    demo_val = 'TRUE' if is_demo_request(event) else 'FALSE'
    where += " AND (p.id IS NULL OR p.is_demo_data = %s)" % demo_val
    conn = get_db()
    cur = conn.cursor()
    if params.get('control_point_id') != 'all':
        where += " AND " + point_sql(load_point(cur, params.get('control_point_id')), 'vp')
    cur.execute("SELECT COUNT(*) FROM vehicle_passes vp LEFT JOIN personnel p ON p.id = vp.personnel_id %s" % where)
    total = cur.fetchone()[0]
    cur.execute("""
        SELECT vp.id, vp.vehicle_number, vp.vehicle_brand, vp.owner_name, vp.direction,
               vp.checkpoint_name, vp.allowed, vp.notes, vp.created_at
        FROM vehicle_passes vp LEFT JOIN personnel p ON p.id = vp.personnel_id
        %s ORDER BY vp.created_at DESC, vp.id DESC LIMIT %d OFFSET %d
    """ % (where, per_page, (page - 1) * per_page))
    rows = cur.fetchall()
    cur.execute("""
        SELECT COUNT(*) FROM (
            SELECT DISTINCT ON (vehicle_number) vehicle_number, direction FROM vehicle_passes
            WHERE allowed = TRUE ORDER BY vehicle_number, created_at DESC, id DESC
        ) t WHERE t.direction = 'in'
    """)
    on_site = cur.fetchone()[0]
    cur.close()
    conn.close()
    items = [{
        'id': r[0], 'vehicle_number': r[1], 'vehicle_brand': r[2], 'owner_name': r[3],
        'direction': 'Въезд' if r[4] == 'in' else 'Выезд', 'direction_raw': r[4],
        'checkpoint_name': r[5], 'allowed': r[6], 'notes': r[7], 'created_at': r[8]
    } for r in rows]
    return json_response(200, {'items': items, 'total': total, 'page': page,
                               'pages': max(1, (total + per_page - 1) // per_page), 'on_site': on_site})


def vehicle_export(params, event):
    """Выгрузка журнала автомобилей в Excel с фильтром по датам"""
    import base64
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    date_from = params.get('date_from', '')
    date_to = params.get('date_to', '')
    try:
        if date_from:
            datetime.strptime(date_from, '%Y-%m-%d')
        if date_to:
            datetime.strptime(date_to, '%Y-%m-%d')
    except ValueError:
        return json_response(400, {'error': 'Неверный формат даты'})
    demo_val = 'TRUE' if is_demo_request(event) else 'FALSE'
    where = "WHERE (p.id IS NULL OR p.is_demo_data = %s)" % demo_val
    if date_from:
        where += " AND vp.created_at >= '%s 00:00:00'" % date_from
    if date_to:
        where += " AND vp.created_at <= '%s 23:59:59'" % date_to
    conn = get_db()
    cur = conn.cursor()
    if params.get('control_point_id') != 'all':
        where += " AND " + point_sql(load_point(cur, params.get('control_point_id')), 'vp')
    cur.execute("""
        SELECT vp.created_at, vp.vehicle_number, vp.vehicle_brand, vp.owner_name,
               COALESCE(p.position, ''), COALESCE(p.organization, ''),
               vp.direction, vp.allowed, vp.notes, vp.checkpoint_name
        FROM vehicle_passes vp LEFT JOIN personnel p ON p.id = vp.personnel_id
        %s ORDER BY vp.created_at DESC, vp.id DESC LIMIT 20000
    """ % where)
    rows = cur.fetchall()
    cur.close()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Журнал автомобилей'
    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    period = 'за всё время'
    if date_from or date_to:
        period = 'с %s по %s' % (
            datetime.strptime(date_from, '%Y-%m-%d').strftime('%d.%m.%Y') if date_from else '…',
            datetime.strptime(date_to, '%Y-%m-%d').strftime('%d.%m.%Y') if date_to else '…')
    ws.cell(row=1, column=1, value='Журнал въезда и выезда автомобилей %s' % period).font = Font(name='Arial', bold=True, size=13)
    allowed_cnt = sum(1 for r in rows if r[7])
    ws.cell(row=2, column=1, value='Всего записей: %d · въездов: %d · выездов: %d · отказов: %d' % (
        len(rows), sum(1 for r in rows if r[7] and r[6] == 'in'), sum(1 for r in rows if r[7] and r[6] == 'out'),
        len(rows) - allowed_cnt)).font = Font(name='Arial', size=10, color='666666')

    headers = ['Дата', 'Время', 'Госномер', 'Марка, модель', 'Владелец', 'Должность', 'Организация', 'Направление', 'Результат', 'Причина отказа', 'КПП']
    widths = [12, 8, 14, 20, 32, 24, 28, 12, 12, 40, 10]
    head_fill = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid')
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
        c.fill = head_fill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    denied_fill = PatternFill(start_color='FDE2E2', end_color='FDE2E2', fill_type='solid')
    for n, r in enumerate(rows, 5):
        reason = (r[8] or '').replace('ОТКАЗ:', '').strip() if not r[7] else ''
        vals = [r[0].strftime('%d.%m.%Y'), r[0].strftime('%H:%M'), r[1], r[2] or '', r[3] or 'Не найден в базе',
                r[4], r[5], 'Въезд' if r[6] == 'in' else 'Выезд', 'Разрешён' if r[7] else 'Отказ', reason, r[9] or '']
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=n, column=i, value=v)
            c.font = Font(name='Arial', size=10, bold=(i == 3), color=('B91C1C' if not r[7] and i in (9, 10) else '000000'))
            c.border = border
            c.alignment = Alignment(vertical='center', wrap_text=(i == 10))
            if not r[7]:
                c.fill = denied_fill
    ws.freeze_panes = 'A5'
    if rows:
        ws.auto_filter.ref = 'A4:K%d' % (len(rows) + 4)

    buf = BytesIO()
    wb.save(buf)
    name = 'Журнал_автомобилей'
    if date_from or date_to:
        name += '_%s_%s' % (date_from or 'начало', date_to or 'сегодня')
    return json_response(200, {'file': base64.b64encode(buf.getvalue()).decode(), 'file_name': name + '.xlsx', 'total': len(rows)})


def detect_housing_conflict(cur, pid):
    """Вернулся хозяин места, где временно живёт другой — фиксирует конфликт расселения"""
    cur.execute("SELECT residence_room_id, full_name FROM personnel WHERE id = %d" % int(pid))
    me = cur.fetchone()
    if not me or not me[0]:
        return []
    rid = me[0]
    cur.execute("""SELECT id, full_name FROM personnel WHERE temp_owner_id = %d AND residence_room_id = %d
                   AND status IN ('arrived', 'on_shift') AND COALESCE(is_hidden, FALSE) = FALSE""" % (int(pid), rid))
    temps = cur.fetchall()
    if not temps:
        return []
    cur.execute("""SELECT r.capacity, r.room_number, COALESCE(b.name, ''),
                   (SELECT COUNT(*) FROM personnel p WHERE p.residence_room_id = r.id AND p.status IN ('arrived', 'on_shift')
                    AND COALESCE(p.is_hidden, FALSE) = FALSE)
                   FROM rooms r LEFT JOIN buildings b ON b.id = r.building_id WHERE r.id = %d""" % rid)
    room = cur.fetchone()
    if not room or room[3] <= room[0]:
        return []
    out = []
    for t in temps:
        cur.execute("SELECT id FROM housing_conflicts WHERE owner_id = %d AND temp_id = %d AND resolved_at IS NULL" % (int(pid), t[0]))
        ex = cur.fetchone()
        if ex:
            cid = ex[0]
        else:
            cur.execute("INSERT INTO housing_conflicts (room_id, owner_id, temp_id) VALUES (%d, %d, %d) RETURNING id" % (rid, int(pid), t[0]))
            cid = cur.fetchone()[0]
            desc = 'Конфликт расселения: вернулся %s (%s, комн. %s), на его месте временно живёт %s' % (me[1], room[2], room[1], t[1])
            cur.execute("INSERT INTO events (event_type, description, personnel_id) VALUES ('housing_conflict', '%s', %d)"
                        % (desc.replace("'", "''"), int(pid)))
        out.append({'id': cid, 'owner_name': me[1], 'temp_name': t[1], 'room_number': room[1], 'building_name': room[2]})
    return out


def parse_qr_code(raw):
    try:
        data = json.loads(raw)
        return data.get('code', raw)
    except (json.JSONDecodeError, AttributeError):
        return raw.strip()

def deny_pass(cur, conn, point, person_id, safe_code, safe_name, person_name, checkpoint_name, medical_ok, reason, direction, person_info):
    cur.execute("""
        INSERT INTO checkpoint_passes (personnel_id, personal_code, full_name, direction, checkpoint_name, medical_ok, notes, control_point_id)
        VALUES (%d, '%s', '%s', '%s', '%s', %s, '%s', %s)
    """ % (person_id, safe_code, safe_name, direction, checkpoint_name.replace("'", "''"),
           'TRUE' if medical_ok else 'FALSE', ('ОТКАЗ: ' + reason).replace("'", "''"), point['id'] if point else 'NULL'))
    cur.execute("""
        INSERT INTO events (event_type, description, personnel_id)
        VALUES ('checkpoint_denied', '%s', %d)
    """ % (('%s: %s — отказ во %s. %s' % (checkpoint_name, person_name, 'входе' if direction == 'in' else 'выходе', reason)).replace("'", "''"), person_id))
    conn.commit()
    cur.close()
    conn.close()
    return json_response(200, {
        'result': 'denied', 'message': 'Проход запрещён', 'reason': reason, 'direction': direction,
        'person_name': person_name, 'medical_ok': medical_ok, 'person': person_info,
        'point': point
    })


def person_on_territory(cur, person_id):
    cur.execute("""SELECT direction FROM checkpoint_passes cp WHERE cp.personnel_id = %d AND %s AND %s
                   ORDER BY cp.created_at DESC, cp.id DESC LIMIT 1""" % (person_id, DENY_SQL % 'cp', territory_sql('cp')))
    r = cur.fetchone()
    return bool(r and r[0] == 'in')


def register_pass(body, event=None):
    raw_code = body.get('code', '').strip()
    direction = body.get('direction', 'in')
    notes = body.get('notes', '')

    if not raw_code:
        return json_response(400, {'error': 'Код не указан'})
    if direction not in ('in', 'out'):
        return json_response(400, {'error': 'direction должен быть in или out'})

    code = parse_qr_code(raw_code)
    safe_code = code.replace("'", "''")

    conn = get_db()
    cur = conn.cursor()
    point = load_point(cur, body.get('control_point_id'))
    if body.get('control_point_id') and not point:
        cur.close()
        conn.close()
        return json_response(404, {'error': 'Объект контроля не найден или отключён'})
    checkpoint_name = str((point['short_name'] if point else '') or body.get('checkpoint_name') or 'КПП-1')[:100]
    settings = point['settings'] if point else DEFAULT_POINT_SETTINGS
    territory = is_territory_point(point)

    cur.execute("""
        SELECT p.id, p.personal_code, p.full_name, p.position, p.department,
               p.category, p.status, p.medical_status, p.organization, p.organization_type,
               COALESCE(p.pass_number, ''), p.pass_permanent, p.pass_valid_from, p.pass_valid_to
        FROM personnel p
        WHERE (p.personal_code = '%s' OR p.qr_code = '%s') AND p.is_hidden = FALSE AND p.status != 'archived'
        LIMIT 1
    """ % (safe_code, safe_code))
    row = cur.fetchone()

    if not row:
        cur.close()
        conn.close()
        return json_response(404, {'error': 'Сотрудник с кодом %s не найден' % code})

    person_id = row[0]
    person_name = row[2]
    safe_name = person_name.replace("'", "''")
    medical_ok = row[7] in ('passed', 'expiring')

    category_labels = {
        'mine': 'Рудничный', 'contractor': 'Подрядчик',
        'business_trip': 'Командированный', 'guest': 'Гость'
    }
    org_type_labels = {
        'rudnik': 'Рудник', 'guest': 'Гость',
        'contractor': 'Подрядная организация', 'gov': 'Гос.органы'
    }
    person_info = {
        'id': row[0],
        'personal_code': row[1],
        'full_name': row[2],
        'position': row[3] or '—',
        'department': row[4] or '—',
        'category': category_labels.get(row[5], row[5]),
        'organization': row[8] or '—',
        'organization_type': org_type_labels.get(row[9] or '', row[9] or '—'),
        'pass_number': row[10] or '',
        'pass_period': pass_period(row[11], row[12], row[13])
    }

    deny = lambda reason: deny_pass(cur, conn, point, person_id, safe_code, safe_name, person_name,
                                    checkpoint_name, medical_ok, reason, direction, person_info)
    if direction == 'in':
        if settings.get('check_pass', True):
            state, reason = check_pass(row[10], row[11], row[12], row[13], 'Личный пропуск')
            if state != 'active':
                return deny(reason)
        if settings.get('require_medical') and not medical_ok:
            return deny('Не пройден медосмотр')
        if not territory and settings.get('require_on_site') and not person_on_territory(cur, person_id):
            return deny('Работник не отмечен на входе через главный КПП')

    cur.execute("""
        INSERT INTO checkpoint_passes (personnel_id, personal_code, full_name, direction, checkpoint_name, medical_ok, notes, control_point_id)
        VALUES (%d, '%s', '%s', '%s', '%s', %s, '%s', %s)
        RETURNING id
    """ % (person_id, safe_code, safe_name, direction, checkpoint_name.replace("'", "''"),
           'TRUE' if medical_ok else 'FALSE', notes.replace("'", "''"), point['id'] if point else 'NULL'))
    pass_id = cur.fetchone()[0]

    housing_conflicts = []
    if not territory:
        event_type = 'point_in' if direction == 'in' else 'point_out'
        event_desc = '%s: %s — %s' % (checkpoint_name.replace("'", "''"), safe_name, 'вход' if direction == 'in' else 'выход')
    elif direction == 'in':
        event_type = 'checkpoint_in'
        event_desc = 'КПП: %s — вход через %s' % (safe_name, checkpoint_name.replace("'", "''"))

        cur.execute("""
            UPDATE personnel SET status = 'arrived', updated_at = NOW() WHERE id = %d
        """ % person_id)

        cur.execute("""
            UPDATE aho_arrivals SET arrival_status = 'arrived', check_in_at = NOW(), updated_at = NOW()
            WHERE personnel_id = %d AND arrival_status = 'expected' AND is_hidden = FALSE
        """ % person_id)
        sync_rooms(cur)
        housing_conflicts = detect_housing_conflict(cur, person_id)

    else:
        event_type = 'checkpoint_out'
        event_desc = 'КПП: %s — выход через %s' % (safe_name, checkpoint_name.replace("'", "''"))

        cur.execute("""
            UPDATE personnel SET status = 'departed', updated_at = NOW() WHERE id = %d
        """ % person_id)

        cur.execute("""
            UPDATE aho_arrivals SET arrival_status = 'departed', check_out_at = NOW(), updated_at = NOW()
            WHERE personnel_id = %d AND arrival_status = 'arrived' AND is_hidden = FALSE
        """ % person_id)

    cur.execute("""
        INSERT INTO events (event_type, description, personnel_id)
        VALUES ('%s', '%s', %d)
    """ % (event_type, event_desc, person_id))
    if territory:
        sync_rooms(cur)

    conn.commit()
    cur.close()
    conn.close()

    dir_label = 'Вход зафиксирован' if direction == 'in' else 'Выход зафиксирован'

    return json_response(200, {
        'result': 'allowed',
        'pass_id': pass_id,
        'message': dir_label,
        'direction': direction,
        'person_name': person_name,
        'medical_ok': medical_ok,
        'person': person_info,
        'housing_conflicts': housing_conflicts,
        'point': point
    })

def get_journal(params, event):
    date_from = params.get('date_from', '')
    date_to = params.get('date_to', '')
    direction = params.get('direction', '')
    page = int(params.get('page', '1'))
    per_page = int(params.get('per_page', '50'))
    offset = (page - 1) * per_page

    conn = get_db()
    cur = conn.cursor()

    demo_val = 'TRUE' if is_demo_request(event) else 'FALSE'
    where = "WHERE (p.id IS NULL OR p.is_demo_data = %s)" % demo_val
    if params.get('control_point_id') != 'all':
        where += " AND " + point_sql(load_point(cur, params.get('control_point_id')), 'cp')
    if date_from:
        where += " AND cp.created_at >= '%s 00:00:00'" % date_from.replace("'", "''")
    if date_to:
        where += " AND cp.created_at <= '%s 23:59:59'" % date_to.replace("'", "''")
    if direction:
        where += " AND cp.direction = '%s'" % direction.replace("'", "''")

    cur.execute("SELECT COUNT(*) FROM checkpoint_passes cp LEFT JOIN personnel p ON cp.personnel_id = p.id %s" % where)
    total = cur.fetchone()[0]

    cur.execute("""
        SELECT cp.id, cp.personnel_id, cp.personal_code, cp.full_name,
               cp.direction, cp.checkpoint_name, cp.medical_ok, cp.notes, cp.created_at,
               COALESCE(p.tab_number, '')
        FROM checkpoint_passes cp
        LEFT JOIN personnel p ON cp.personnel_id = p.id
        %s
        ORDER BY cp.created_at DESC
        LIMIT %d OFFSET %d
    """ % (where, per_page, offset))
    rows = cur.fetchall()

    cur.close()
    conn.close()

    items = [{
        'id': r[0], 'personnel_id': r[1], 'personal_code': r[2],
        'full_name': r[3],
        'direction': 'Вход' if r[4] == 'in' else 'Выход',
        'direction_raw': r[4],
        'checkpoint_name': r[5],
        'medical_ok': r[6],
        'notes': r[7] or '',
        'created_at': r[8],
        'tab_number': r[9] or ''
    } for r in rows]

    return json_response(200, {
        'items': items, 'total': total,
        'page': page, 'per_page': per_page,
        'pages': (total + per_page - 1) // per_page if per_page > 0 else 1
    })

def get_stats(event, params=None):
    params = params or {}
    conn = get_db()
    cur = conn.cursor()

    demo_val = 'TRUE' if is_demo_request(event) else 'FALSE'
    point = load_point(cur, params.get('control_point_id'))
    flt = point_sql(point, 'cp')
    cur.execute("""
        SELECT COUNT(*) FILTER (WHERE cp.direction = 'in' AND %s AND cp.created_at::date = CURRENT_DATE),
               COUNT(*) FILTER (WHERE cp.direction = 'out' AND %s AND cp.created_at::date = CURRENT_DATE),
               COUNT(*) FILTER (WHERE NOT (%s) AND cp.created_at::date = CURRENT_DATE),
               COUNT(*)
        FROM checkpoint_passes cp JOIN personnel p ON cp.personnel_id = p.id AND p.is_demo_data = %s
        WHERE %s
    """ % (DENY_SQL % 'cp', DENY_SQL % 'cp', DENY_SQL % 'cp', demo_val, flt))
    t = cur.fetchone()
    cur.execute("SELECT COUNT(*) FROM (%s) x" % inside_sql(point, demo_val))
    currently_on_site = cur.fetchone()[0]

    cur.close()
    conn.close()

    return json_response(200, {
        'today_in': t[0],
        'today_out': t[1],
        'today_denied': t[2],
        'total_passes': t[3],
        'currently_on_site': currently_on_site,
        'point': point
    })

def get_on_site(params, event):
    conn = get_db()
    cur = conn.cursor()

    demo_val = 'TRUE' if is_demo_request(event) else 'FALSE'
    point = load_point(cur, params.get('control_point_id'))
    cur.execute("""
        SELECT ins.personnel_id, p.personal_code, p.full_name, ins.checkpoint_name, ins.created_at,
               p.position, p.department, p.organization, p.organization_type, COALESCE(p.tab_number, '')
        FROM (%s) ins JOIN personnel p ON p.id = ins.personnel_id
        ORDER BY ins.created_at DESC
    """ % inside_sql(point, demo_val))
    rows = cur.fetchall()
    cur.close()
    conn.close()

    org_type_labels = {
        'rudnik': 'Рудник', 'guest': 'Гость',
        'contractor': 'Подрядная организация', 'gov': 'Гос.органы'
    }
    on_site = [{
        'personnel_id': r[0], 'personal_code': r[1], 'full_name': r[2],
        'checkpoint_name': r[3], 'entered_at': r[4],
        'position': r[5] or '—', 'department': r[6] or '—',
        'organization': r[7] or '—',
        'organization_type': org_type_labels.get(r[8] or '', r[8] or '—'),
        'tab_number': r[9] or ''
    } for r in rows]

    return json_response(200, {'items': on_site, 'total': len(on_site)})

def export_journal(params, event):
    date_from = params.get('date_from', '')
    date_to = params.get('date_to', '')

    conn = get_db()
    cur = conn.cursor()

    demo_val = 'TRUE' if is_demo_request(event) else 'FALSE'
    where = "WHERE (p.id IS NULL OR p.is_demo_data = %s)" % demo_val
    if params.get('control_point_id') != 'all':
        where += " AND " + point_sql(load_point(cur, params.get('control_point_id')), 'cp')
    if date_from:
        where += " AND cp.created_at >= '%s 00:00:00'" % date_from.replace("'", "''")
    if date_to:
        where += " AND cp.created_at <= '%s 23:59:59'" % date_to.replace("'", "''")

    cur.execute("""
        SELECT cp.personal_code, cp.full_name, cp.direction, cp.checkpoint_name,
               cp.medical_ok, cp.notes, cp.created_at, COALESCE(p.tab_number, '')
        FROM checkpoint_passes cp
        LEFT JOIN personnel p ON cp.personnel_id = p.id
        %s
        ORDER BY cp.created_at DESC
        LIMIT 5000
    """ % where)
    rows = cur.fetchall()
    cur.close()
    conn.close()

    show_code = caller_is_admin(event)
    lines = [('Код;' if show_code else '') + 'ФИО;Таб. №;Направление;Объект;Медосмотр;Примечание;Дата и время']
    for r in rows:
        dt = r[6].strftime('%d.%m.%Y %H:%M') if r[6] else ''
        direction = 'Вход' if r[2] == 'in' else 'Выход'
        medical = 'Да' if r[4] else 'Нет'
        lines.append(('%s;' % r[0] if show_code else '') + '%s;%s;%s;%s;%s;%s;%s' % (
            r[1], r[7] or '', direction, r[3], medical, r[5] or '', dt
        ))

    csv_content = '\n'.join(lines)

    return {
        'statusCode': 200,
        'headers': {
            'Content-Type': 'text/csv; charset=utf-8',
            'Content-Disposition': 'attachment; filename="checkpoint_journal.csv"',
            'Access-Control-Allow-Origin': '*'
        },
        'body': csv_content
    }

def get_auth_token(event):
    headers = event.get('headers') or {}
    for key in ['X-Authorization', 'x-authorization', 'Authorization', 'authorization']:
        val = headers.get(key, '')
        if val:
            return val[7:] if val.startswith('Bearer ') else val
    return ''


def get_admin(cur, event):
    token = get_auth_token(event)
    if not token:
        return None
    cur.execute("""
        SELECT u.id, u.full_name FROM users u
        JOIN sessions s ON s.user_id = u.id
        WHERE s.token = '%s' AND s.expires_at > NOW() AND u.is_active = TRUE AND u.role = 'admin'
    """ % token.replace("'", "''"))
    return cur.fetchone()


def sync_person_status(cur, personnel_id):
    if not personnel_id:
        return
    cur.execute("""
        SELECT cp.direction FROM checkpoint_passes cp
        WHERE cp.personnel_id = %d AND %s AND %s
        ORDER BY cp.created_at DESC, cp.id DESC LIMIT 1
    """ % (int(personnel_id), DENY_SQL % 'cp', territory_sql('cp')))
    row = cur.fetchone()
    if not row:
        return
    new_status = 'arrived' if row[0] == 'in' else 'departed'
    cur.execute("""
        UPDATE personnel SET status = '%s', updated_at = NOW()
        WHERE id = %d AND status IN ('arrived', 'departed', 'on_shift')
    """ % (new_status, int(personnel_id)))
    sync_rooms(cur)


def update_pass(event, body):
    pass_id = body.get('id')
    if not pass_id:
        return json_response(400, {'error': 'Не указана запись'})

    conn = get_db()
    cur = conn.cursor()
    admin = get_admin(cur, event)
    if not admin:
        cur.close()
        conn.close()
        return json_response(403, {'error': 'Редактировать журнал может только администратор'})

    cur.execute("SELECT id, personnel_id, full_name FROM checkpoint_passes WHERE id = %d" % int(pass_id))
    row = cur.fetchone()
    if not row:
        cur.close()
        conn.close()
        return json_response(404, {'error': 'Запись не найдена'})

    sets = []
    direction = body.get('direction')
    if direction is not None:
        if direction not in ('in', 'out'):
            cur.close()
            conn.close()
            return json_response(400, {'error': 'Неверное направление'})
        sets.append("direction = '%s'" % direction)
    if body.get('checkpoint_name') is not None:
        name = str(body['checkpoint_name']).strip()[:100] or 'КПП-1'
        sets.append("checkpoint_name = '%s'" % name.replace("'", "''"))
    if body.get('notes') is not None:
        sets.append("notes = '%s'" % str(body['notes']).replace("'", "''"))
    if body.get('medical_ok') is not None:
        sets.append("medical_ok = %s" % ('TRUE' if body['medical_ok'] else 'FALSE'))
    if body.get('created_at'):
        raw = str(body['created_at']).strip().replace('Z', '+00:00')
        try:
            dt = datetime.fromisoformat(raw)
        except ValueError:
            cur.close()
            conn.close()
            return json_response(400, {'error': 'Неверная дата'})
        if dt.tzinfo is not None:
            from datetime import timezone
            dt = dt.astimezone(timezone.utc).replace(tzinfo=None)
        sets.append("created_at = '%s'" % dt.strftime('%Y-%m-%d %H:%M:%S'))

    if not sets:
        cur.close()
        conn.close()
        return json_response(400, {'error': 'Нет изменений'})

    cur.execute("UPDATE checkpoint_passes SET %s WHERE id = %d" % (', '.join(sets), int(pass_id)))
    sync_person_status(cur, row[1])
    conn.commit()
    cur.close()
    conn.close()
    return json_response(200, {'success': True})


def delete_pass(event, body, params):
    pass_id = body.get('id') or params.get('id')
    if not pass_id:
        return json_response(400, {'error': 'Не указана запись'})

    conn = get_db()
    cur = conn.cursor()
    admin = get_admin(cur, event)
    if not admin:
        cur.close()
        conn.close()
        return json_response(403, {'error': 'Удалять записи журнала может только администратор'})

    cur.execute("SELECT id, personnel_id, full_name, direction, created_at FROM checkpoint_passes WHERE id = %d" % int(pass_id))
    row = cur.fetchone()
    if not row:
        cur.close()
        conn.close()
        return json_response(404, {'error': 'Запись не найдена'})

    cur.execute("DELETE FROM checkpoint_passes WHERE id = %d" % int(pass_id))
    sync_person_status(cur, row[1])
    conn.commit()
    cur.close()
    conn.close()
    return json_response(200, {'success': True})
